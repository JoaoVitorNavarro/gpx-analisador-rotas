"""Exportacao da planilha Excel consolidada (openpyxl).

Gera um unico arquivo com as abas: Resumo, Ruas percorridas, Diagnostico dos
pontos (opcional) e Erros e alertas. Formatacao profissional: cabecalhos
destacados, congelamento da 1a linha, autofiltro, larguras ajustadas, formatos
de data/hora/distancia e cores por nivel de confianca. Nunca sobrescreve
arquivos existentes.

Observacao sobre datas: o Excel nao armazena fuso horario; os horarios do GPX
sao UTC e sao gravados como data/hora "ingenua" em UTC (indicado no cabecalho).
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import utils
from .config_loader import Config
from .models import (CONF_ALTA, CONF_BAIXA, CONF_MEDIA, CONF_NAO_IDENT, FileResult)

logger = utils.get_logger()

# estilos
_HDR_FILL = PatternFill("solid", fgColor="1F4E78")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BORDA = Border(*(Side(style="thin", color="D9D9D9"),) * 4)

_COR_CONF = {
    CONF_ALTA: PatternFill("solid", fgColor="C6EFCE"),
    CONF_MEDIA: PatternFill("solid", fgColor="FFEB9C"),
    CONF_BAIXA: PatternFill("solid", fgColor="FFC7A8"),
    CONF_NAO_IDENT: PatternFill("solid", fgColor="FFC7CE"),
}
_COR_STATUS = {
    "ok": PatternFill("solid", fgColor="C6EFCE"),
    "parcial": PatternFill("solid", fgColor="FFEB9C"),
    "erro": PatternFill("solid", fgColor="FFC7CE"),
    "cancelado": PatternFill("solid", fgColor="D9D9D9"),
}

_FMT_DIST = "#,##0.0"
_FMT_DIST2 = "#,##0.00"
_FMT_DT = "yyyy-mm-dd hh:mm:ss"
_FMT_LATLON = "0.000000"


def _naive(dt: _dt.datetime | None):
    if dt is None:
        return None
    return dt.replace(tzinfo=None)


def _cabecalho(ws, colunas: list[str]) -> None:
    for c, titulo in enumerate(colunas, start=1):
        cell = ws.cell(row=1, column=c, value=titulo)
        cell.fill = _HDR_FILL
        cell.font = _HDR_FONT
        cell.alignment = _HDR_ALIGN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}1"


def _ajustar_larguras(ws, larguras: list[int]) -> None:
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def exportar(resultados: list[FileResult], cfg: Config, saida_dir: Path) -> Path:
    saida_dir = Path(saida_dir)
    saida_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    _aba_resumo(wb.active, resultados)
    _aba_ruas(wb.create_sheet("Ruas percorridas"), resultados)
    if bool(cfg.get("gerar_diagnostico_pontos", False)):
        _aba_diagnostico(wb.create_sheet("Diagnostico dos pontos"), resultados)
    _aba_erros(wb.create_sheet("Erros e alertas"), resultados)

    ts = _dt.datetime.now().strftime("%d%m%Y_%H%M%S")
    destino = utils.caminho_unico(saida_dir / f"Relatorio_Rotas_GPX_{ts}.xlsx")
    wb.save(destino)
    logger.info("Excel gerado: %s", destino)
    return destino


# --------------------------------------------------------------------------- Resumo
def _aba_resumo(ws, resultados: list[FileResult]) -> None:
    ws.title = "Resumo"
    colunas = [
        "Arquivo GPX", "Identificador do video", "Data inicial (UTC)",
        "Data final (UTC)", "Duracao", "Distancia total (m)", "Pontos originais",
        "Pontos apos limpeza", "Pontos processados", "Qtd. de vias",
        "Confianca alta", "Confianca media", "Confianca baixa",
        "Nao identificados", "Situacao", "Caminho do mapa", "Observacoes",
    ]
    _cabecalho(ws, colunas)
    r = 2
    for fr in resultados:
        info = fr.info
        resumo = fr.resumo_confianca()
        obs = "; ".join(info.anomalias) if info and info.anomalias else ""
        vals = [
            fr.nome, fr.id_video,
            _naive(info.tempo_inicial) if info else None,
            _naive(info.tempo_final) if info else None,
            utils.formatar_duracao(info.duracao_s if info else None),
            round(info.distancia_total_m, 1) if info else 0.0,
            fr.simplify.n_original if fr.simplify else 0,
            fr.simplify.n_apos_limpeza if fr.simplify else 0,
            fr.simplify.n_apos_simplificacao if fr.simplify else 0,
            len(fr.trechos),
            resumo[CONF_ALTA], resumo[CONF_MEDIA], resumo[CONF_BAIXA], resumo[CONF_NAO_IDENT],
            fr.status, fr.caminho_mapa, obs,
        ]
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v).border = _BORDA
        ws.cell(row=r, column=3).number_format = _FMT_DT
        ws.cell(row=r, column=4).number_format = _FMT_DT
        ws.cell(row=r, column=6).number_format = _FMT_DIST
        cell_status = ws.cell(row=r, column=15)
        if fr.status in _COR_STATUS:
            cell_status.fill = _COR_STATUS[fr.status]
        r += 1
    _ajustar_larguras(ws, [26, 16, 20, 20, 10, 16, 13, 15, 15, 10, 12, 12, 12, 14, 10, 40, 40])


# --------------------------------------------------------------------------- Ruas
def _aba_ruas(ws, resultados: list[FileResult]) -> None:
    colunas = [
        "Arquivo GPX", "Identificador do video", "Ordem", "Nome da via",
        "Referencia", "Classe da via", "OSM ID", "Horario de entrada (UTC)",
        "Horario de saida (UTC)", "Duracao", "Distancia percorrida (m)",
        "Lat inicial", "Lon inicial", "Lat final", "Lon final",
        "Qtd. de pontos", "Dist. media ate a via (m)", "Dist. maxima (m)",
        "Confianca", "Observacoes",
    ]
    _cabecalho(ws, colunas)
    r = 2
    for fr in resultados:
        for t in fr.trechos:
            vals = [
                fr.nome, fr.id_video, t.ordem, t.nome_via, t.ref, t.classe,
                str(t.osm_id) if t.osm_id is not None else "",
                _naive(t.hora_entrada), _naive(t.hora_saida),
                utils.formatar_duracao(t.duracao_s),
                round(t.distancia_m, 1), t.lat_ini, t.lon_ini, t.lat_fim, t.lon_fim,
                t.n_pontos, round(t.dist_media_m, 2), round(t.dist_max_m, 2),
                t.confianca, t.observacoes,
            ]
            for c, v in enumerate(vals, start=1):
                ws.cell(row=r, column=c, value=v).border = _BORDA
            ws.cell(row=r, column=8).number_format = _FMT_DT
            ws.cell(row=r, column=9).number_format = _FMT_DT
            ws.cell(row=r, column=11).number_format = _FMT_DIST
            for col in (12, 13, 14, 15):
                ws.cell(row=r, column=col).number_format = _FMT_LATLON
            for col in (17, 18):
                ws.cell(row=r, column=col).number_format = _FMT_DIST2
            cell_conf = ws.cell(row=r, column=19)
            if t.confianca in _COR_CONF:
                cell_conf.fill = _COR_CONF[t.confianca]
            r += 1
    _ajustar_larguras(ws, [24, 15, 7, 28, 12, 14, 14, 20, 20, 10, 18, 12, 12, 12, 12,
                           14, 20, 16, 14, 36])


# --------------------------------------------------------------------------- Diagnostico
def _aba_diagnostico(ws, resultados: list[FileResult]) -> None:
    colunas = [
        "Arquivo", "Ordem (simplificado)", "Data e hora (UTC)", "Latitude",
        "Longitude", "Distancia acumulada (m)", "Velocidade (km/h)", "Direcao (graus)",
        "Via selecionada", "OSM ID", "Distancia ate a via (m)",
        "Segunda melhor candidata", "Diferenca de pontuacao", "Confianca",
        "Motivo de descarte",
    ]
    _cabecalho(ws, colunas)
    r = 2
    for fr in resultados:
        conf_por_idx: dict[int, str] = {}
        for t in fr.trechos:
            for idx in t.indices:
                conf_por_idx[idx] = t.confianca
        for mp in fr.matched:
            esc = mp.escolhido
            seg = mp.segundo
            vals = [
                fr.nome, mp.indice, _naive(mp.tp.tempo), mp.tp.lat, mp.tp.lon,
                round(mp.tp.dist_acum_m, 1),
                round(mp.tp.vel_kmh, 1) if mp.tp.vel_kmh is not None else None,
                round(mp.tp.azimute, 1) if mp.tp.azimute is not None else None,
                esc.nome if esc else "", str(esc.osmid) if esc and esc.osmid is not None else "",
                round(mp.dist_via_m, 2) if mp.dist_via_m is not None else None,
                seg.nome if seg else "",
                round(mp.margem_score, 3) if mp.margem_score is not None else None,
                conf_por_idx.get(mp.indice, ""), mp.motivo_descarte,
            ]
            for c, v in enumerate(vals, start=1):
                ws.cell(row=r, column=c, value=v).border = _BORDA
            ws.cell(row=r, column=3).number_format = _FMT_DT
            ws.cell(row=r, column=4).number_format = _FMT_LATLON
            ws.cell(row=r, column=5).number_format = _FMT_LATLON
            r += 1
    _ajustar_larguras(ws, [22, 18, 20, 12, 12, 18, 14, 14, 26, 14, 18, 26, 18, 12, 26])


# --------------------------------------------------------------------------- Erros
def _aba_erros(ws, resultados: list[FileResult]) -> None:
    colunas = ["Arquivo", "Etapa", "Tipo", "Mensagem", "Data e hora", "Acao recomendada"]
    _cabecalho(ws, colunas)
    r = 2
    algum = False
    for fr in resultados:
        for a in fr.alertas:
            algum = True
            vals = [a.arquivo, a.etapa, a.tipo, a.mensagem,
                    a.data_hora or fr.data_processamento, a.acao]
            for c, v in enumerate(vals, start=1):
                ws.cell(row=r, column=c, value=v).border = _BORDA
            if a.tipo == "ERRO":
                ws.cell(row=r, column=3).fill = _COR_CONF[CONF_NAO_IDENT]
            elif a.tipo == "ALERTA":
                ws.cell(row=r, column=3).fill = _COR_CONF[CONF_MEDIA]
            r += 1
    if not algum:
        ws.cell(row=2, column=1, value="Nenhum erro ou alerta registrado.")
    _ajustar_larguras(ws, [24, 16, 10, 60, 20, 40])
