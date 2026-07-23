"""Testes de geracao de saidas (Excel) e teste de integracao com o GPX real."""
from __future__ import annotations

import datetime as _dt

import openpyxl
import pytest

from src import config_loader, excel_exporter
from src.models import CONF_ALTA, FileResult, GpxInfo, SimplifyStats, Trecho


def _file_result_sintetico() -> FileResult:
    fr = FileResult(caminho="x/a.gpx", nome="a.gpx", id_video="GX0001", status="ok")
    fr.info = GpxInfo(caminho="x/a.gpx", nome="a.gpx", n_trilhas=1, n_segmentos=1,
                      n_pontos=100, distancia_total_m=500.0,
                      tempo_inicial=_dt.datetime(2026, 7, 1, 12, 0, 0),
                      tempo_final=_dt.datetime(2026, 7, 1, 12, 5, 0), duracao_s=300.0)
    fr.simplify = SimplifyStats(n_original=100, n_apos_limpeza=80, n_apos_simplificacao=30)
    fr.trechos = [
        Trecho(ordem=1, road_id="A", nome_via="Rua das Pratas", classe="residential",
               osm_id=123, n_pontos=20, distancia_m=300.0, dist_media_m=4.0,
               dist_max_m=9.0, confianca=CONF_ALTA,
               hora_entrada=_dt.datetime(2026, 7, 1, 12, 0, 0),
               hora_saida=_dt.datetime(2026, 7, 1, 12, 3, 0),
               lat_ini=-3.02, lon_ini=-59.94, lat_fim=-3.021, lon_fim=-59.941),
    ]
    return fr


def test_excel_gera_abas(tmp_path):
    cfg = config_loader.carregar_config()
    cfg.set("gerar_diagnostico_pontos", False)
    fr = _file_result_sintetico()
    caminho = excel_exporter.exportar([fr], cfg, tmp_path)
    assert caminho.exists()
    wb = openpyxl.load_workbook(caminho)
    assert "Resumo" in wb.sheetnames
    assert "Ruas percorridas" in wb.sheetnames
    assert "Erros e alertas" in wb.sheetnames
    ws = wb["Ruas percorridas"]
    assert ws.cell(2, 4).value == "Rua das Pratas"
    assert ws.freeze_panes == "A2"


def test_excel_nao_sobrescreve(tmp_path):
    cfg = config_loader.carregar_config()
    fr = _file_result_sintetico()
    c1 = excel_exporter.exportar([fr], cfg, tmp_path)
    c2 = excel_exporter.exportar([fr], cfg, tmp_path)
    assert c1 != c2  # nomes com timestamp/sufixo diferentes
    assert c1.exists() and c2.exists()


@pytest.mark.integracao
def test_integracao_gpx_real_offline(cfg, gpx_teste_path, tmp_path):
    """Integracao ponta a ponta com o GPX real (offline, requer cache pronto)."""
    if gpx_teste_path is None:
        pytest.skip("GPX de teste indisponivel neste ambiente")
    from src.database import Database
    from src.route_processor import BatchProcessor, Reporter
    cfg.set("modo_offline", True)
    db = Database()
    try:
        bp = BatchProcessor(cfg, db, Reporter())
        res = bp.processar([str(gpx_teste_path)], tmp_path)
    finally:
        db.fechar()
    assert len(res) == 1
    fr = res[0]
    # se o cache estiver pronto, deve identificar vias; senao, parcial (ainda valido)
    if fr.status == "ok":
        assert len(fr.trechos) >= 1
        assert any(t.road_id for t in fr.trechos)
